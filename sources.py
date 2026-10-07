# -*- coding: utf-8 -*-
# OSInt Graph - esplorazione grafica di dati OSINT per casi investigativi
# Copyright (C) 2025-2026 Andrea Cumini <andrea@osintinfo.net> - www.osintinfo.net
# SPDX-License-Identifier: GPL-3.0-or-later
#
# This program is free software: you can redistribute it and/or modify it
# under the terms of the GNU General Public License as published by the Free
# Software Foundation, either version 3 of the License, or (at your option)
# any later version, with the additional terms in the file NOTICE.
# See the files LICENSE and NOTICE for details. Distributed WITHOUT ANY WARRANTY.

"""Indicizzazione di una cartella di fonti (documenti) e ricerca full-text.

Formati supportati per l'estrazione del testo:
    .txt .md .csv .log    lettura diretta
    .docx                 python-docx
    .doc                  best-effort (stringhe leggibili dal binario)
    .xlsx                 openpyxl
    .xls                  xlrd
    .pdf                  pypdf (PDF testuali)
    .png .jpg .jpeg ...   OCR con Tesseract, SOLO se installato

L'indice lo salva il chiamante (nel database dell'app) così sopravvive ai riavvii.
La ricerca è per sottostringa (parziale) o per parola intera (totale), senza
distinzione di maiuscole/accenti, e restituisce anteprime di contesto.
"""
import os
import re
import io
import json
import time
import hashlib
import logging
import itertools
import unicodedata

TEXT_EXT = {".txt", ".md", ".csv", ".log", ".text", ".json"}
DOCX_EXT = {".docx"}
DOC_EXT = {".doc"}
XLSX_EXT = {".xlsx", ".xlsm"}
XLS_EXT = {".xls"}
PDF_EXT = {".pdf"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
SUPPORTED = TEXT_EXT | DOCX_EXT | DOC_EXT | XLSX_EXT | XLS_EXT | PDF_EXT | IMG_EXT

MAX_TEXT = 800_000          # testo massimo tenuto per file (caratteri)
MAX_FILE_MB = 60            # oltre questa dimensione il file viene saltato


# ---------------------------------------------------------------- OCR ------
# Tesseract può stare in una cartella LOCALE (scaricata con setup_tesseract.py),
# senza installarlo nel sistema. Si cerca lì prima che nel PATH.
def _local_tesseract_candidates(base_dir):
    exe = "tesseract.exe" if os.name == "nt" else "tesseract"
    dirs = []
    if os.environ.get("TESSERACT_DIR"):
        dirs.append(os.environ["TESSERACT_DIR"])
    if base_dir:
        dirs.append(os.path.join(base_dir, "tools", "tesseract"))
    dirs.append(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             "tools", "tesseract"))
    out = []
    for d in dirs:
        out.append(os.path.join(d, exe))
        out.append(os.path.join(d, "bin", exe))       # layout con sottocartella bin/
    return out


TESSDATA_DIR = ""          # cartella tessdata locale, se trovata


OCR_REASON = ""     # perche' l'OCR manca: "module" (pytesseract), "binary" (tesseract) o ""


def configure_ocr(base_dir=None):
    """Configura pytesseract usando un Tesseract locale se presente, altrimenti
    quello di sistema. Aggiorna e ritorna OCR_AVAILABLE (e OCR_REASON)."""
    global OCR_AVAILABLE, TESSDATA_DIR, OCR_REASON
    try:
        import pytesseract
    except Exception:
        OCR_AVAILABLE, OCR_REASON = False, "module"
        return False
    # 1) binario locale (cartella tools/tesseract)
    for cand in _local_tesseract_candidates(base_dir):
        if os.path.isfile(cand):
            pytesseract.pytesseract.tesseract_cmd = cand
            # i dati lingua stanno accanto, in tessdata/. Tesseract 5.x vuole
            # TESSDATA_PREFIX puntato DIRETTAMENTE a quella cartella.
            tdata = os.path.join(os.path.dirname(cand), "tessdata")
            if os.path.isdir(tdata):
                TESSDATA_DIR = tdata
                os.environ["TESSDATA_PREFIX"] = tdata
            break
    # 2) verifica che funzioni (locale o di sistema)
    try:
        pytesseract.get_tesseract_version()
        OCR_AVAILABLE, OCR_REASON = True, ""
    except Exception:
        OCR_AVAILABLE, OCR_REASON = False, "binary"
    return OCR_AVAILABLE


OCR_AVAILABLE = configure_ocr()


# --------------------------------------------------- estrattori di testo ---
def _txt(path):
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with io.open(path, "r", encoding=enc) as f:
                return f.read()
        except Exception:
            continue
    return ""


def _docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for tbl in d.tables:                        # anche il testo nelle tabelle
        for row in tbl.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _doc(path):
    """Vecchio formato .doc: estrae le sequenze di caratteri leggibili dal
    binario. Grezzo ma cattura il testo visibile senza dipendenze esterne."""
    try:
        raw = io.open(path, "rb").read()
    except Exception:
        return ""
    txt = raw.decode("latin-1", "ignore")
    tokens = re.findall(r"[ -~ -ɏ]{4,}", txt)   # ASCII + latino esteso
    return "\n".join(t.strip() for t in tokens if len(t.strip()) >= 4)


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c not in (None, "")]
            if cells:
                out.append("\t".join(cells))
    wb.close()
    return "\n".join(out)


def _xls(path):
    import xlrd
    book = xlrd.open_workbook(path)
    out = []
    for sh in book.sheets():
        out.append(f"# {sh.name}")
        for r in range(sh.nrows):
            cells = [str(c) for c in sh.row_values(r) if str(c).strip()]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)


def _pdf(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    out = []
    for pg in r.pages:
        try:
            out.append(pg.extract_text() or "")
        except Exception:
            pass
        if sum(len(x) for x in out) > MAX_TEXT:
            break
    return "\n".join(out)


def _image(path):
    if not OCR_AVAILABLE:
        return ""
    try:
        import pytesseract
        from PIL import Image
        # la posizione dei dati lingua è data da TESSDATA_PREFIX (impostato in
        # configure_ocr direttamente sulla cartella tessdata, come vuole Tesseract 5).
        return pytesseract.image_to_string(Image.open(path))
    except Exception as e:
        logging.warning(f"OCR fallito su {path}: {e}")
        return ""


def extract_text(path, ext):
    try:
        if ext in TEXT_EXT:
            return _txt(path)
        if ext in DOCX_EXT:
            return _docx(path)
        if ext in DOC_EXT:
            return _doc(path)
        if ext in XLSX_EXT:
            return _xlsx(path)
        if ext in XLS_EXT:
            return _xls(path)
        if ext in PDF_EXT:
            return _pdf(path)
        if ext in IMG_EXT:
            return _image(path)
    except Exception as e:
        logging.warning(f"Estrazione fallita su {path}: {e}")
    return ""


# ---------------------------------------------------- normalizzazione ------
def _norm_char(c):
    """Normalizza UN carattere restando 1:1 (un char -> un char): toglie
    l'accento e mette in minuscolo, senza cambiare la lunghezza. Fondamentale
    perché gli offset dei match combacino col testo originale (highlight)."""
    d = unicodedata.normalize("NFKD", c)
    d = "".join(ch for ch in d if not unicodedata.combining(ch))
    return (d[0] if d else c).lower()


def norm(s):
    """Minuscolo, accenti rimossi, MA con mappatura 1:1 dei caratteri, così la
    posizione di ogni carattere è la stessa del testo originale."""
    return "".join(_norm_char(c) for c in str(s or ""))


def file_id(path):
    return hashlib.md5(os.path.abspath(path).encode("utf-8")).hexdigest()[:16]


# ------------------------------------------------------------- scansione ---
def _list_candidates(folder):
    """Elenco dei file supportati (path, name, ext, size), saltando i troppo
    grandi. Serve a conoscere il TOTALE prima di iniziare a estrarre."""
    out, skipped = [], 0
    for root, dirs, names in os.walk(folder):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for name in names:
            ext = os.path.splitext(name)[1].lower()
            if ext not in SUPPORTED:
                continue
            path = os.path.join(root, name)
            try:
                size = os.path.getsize(path)
            except OSError:
                continue
            if size > MAX_FILE_MB * 1024 * 1024:
                skipped += 1
                continue
            out.append((path, name, ext, size))
    return out, skipped


def _rel_to_base(folder, base_dir):
    """Se 'folder' sta dentro base_dir, ritorna il percorso relativo (per la
    PORTABILITÀ: spostando la cartella, l'indice resta valido). Altrimenti None."""
    if not base_dir:
        return None
    try:
        base = os.path.abspath(base_dir)
        rel = os.path.relpath(folder, base)
    except ValueError:                      # dischi diversi su Windows
        return None
    # dentro base_dir se il relativo non risale con '..' e non è assoluto
    if rel == os.pardir or rel.startswith(os.pardir + os.sep) or os.path.isabs(rel):
        return None
    return rel.replace("\\", "/")


def effective_folder(index, base_dir):
    """Cartella fonti CORRENTE: se l'indice ha un percorso relativo (cartella
    dentro/accanto all'eseguibile) lo risolve rispetto a base_dir ORA — così
    tutto resta valido anche dopo aver spostato la cartella."""
    rel = index.get("folder_rel")
    if rel and base_dir:
        return os.path.abspath(os.path.join(base_dir, rel.replace("/", os.sep)))
    return index.get("folder") or ""


def folder_candidates(index, base_dir, default_folder=None):
    """Cartelle dove cercare OGGI i documenti indicizzati, IN ORDINE.

    Prima si guarda SEMPRE nella cartella 'fonti' accanto all'eseguibile: e'
    quella la posizione ufficiale dei documenti e l'unica che resta valida
    spostando il programma da un PC all'altro.
      1) fonti/<nome della cartella scansionata>   (es. fonti/Tutti_dati_Lybia)
      2) fonti/
      3) il percorso registrato nell'indice (assoluto: vale solo sulla
         macchina che ha fatto la scansione)

    Senza il punto 1-2 succedeva questo: l'indice conserva il percorso della
    cartella scansionata e, se quella stava FUORI dalla directory del
    programma, il percorso e' assoluto. Portando eseguibile + data/ + fonti/
    su un altro PC la ricerca continuava a funzionare (il testo e' dentro
    l'indice) ma l'apertura dei file rispondeva 404."""
    out = []

    def aggiungi(p):
        if not p:
            return
        p = os.path.abspath(p)
        if p not in out:
            out.append(p)

    fonti = default_folder or (os.path.join(base_dir, "fonti") if base_dir else None)
    if fonti:
        originale = (index.get("folder") or "").replace("/", os.sep).rstrip("\/")
        nome = os.path.basename(originale)
        if nome:
            aggiungi(os.path.join(fonti, nome))
        aggiungi(fonti)
        # ultimo tentativo dentro 'fonti': le sottocartelle di primo livello.
        # Copre il caso in cui i documenti siano stati messi li' sotto un nome
        # diverso da quello della cartella scansionata.
        try:
            for voce in sorted(os.listdir(fonti))[:40]:
                sotto = os.path.join(fonti, voce)
                if os.path.isdir(sotto):
                    aggiungi(sotto)
        except OSError:
            pass
    aggiungi(effective_folder(index, base_dir))
    return out


def locate_file(index, rel, base_dir, default_folder=None):
    """Percorso reale di un file indicizzato, cercato fra le cartelle possibili.
    Ritorna (percorso, cartella) oppure (None, None) se non si trova.
    Il controllo di sicurezza resta: il file deve stare DENTRO la cartella."""
    rel = (rel or "").replace("/", os.sep)
    if not rel:
        return None, None
    for cartella in folder_candidates(index, base_dir, default_folder):
        percorso = os.path.abspath(os.path.join(cartella, rel))
        if not percorso.startswith(cartella + os.sep):
            continue                      # rel che tenta di uscire: si scarta
        if os.path.isfile(percorso):
            return percorso, cartella
    return None, None


def resolved_folder(index, base_dir, default_folder=None):
    """Cartella da cui si stanno leggendo DAVVERO i documenti: la prima
    candidata che esiste e contiene almeno un file dell'indice."""
    files = index.get("files") or []
    campione = [f.get("rel") for f in files[:25] if f.get("rel")]
    for cartella in folder_candidates(index, base_dir, default_folder):
        if not os.path.isdir(cartella):
            continue
        if not campione:
            return cartella
        for rel in campione:
            if os.path.isfile(os.path.join(cartella, rel.replace("/", os.sep))):
                return cartella
    return effective_folder(index, base_dir)


def scan_folder_stream(folder, index_path, base_dir=None):
    """Come scan_folder, ma è un GENERATORE che emette l'avanzamento a ogni
    file elaborato: prima {'phase':'start','total':N}, poi per ciascun file
    {'phase':'file','done':i,'total':N,'name':...}, infine {'phase':'done', ...}
    con il riepilogo. Salva l'indice alla fine. Se la cartella sta dentro
    base_dir, memorizza il percorso RELATIVO (portabilità)."""
    folder = os.path.abspath(folder)
    if not os.path.isdir(folder):
        raise FileNotFoundError(folder)

    folder_rel = _rel_to_base(folder, base_dir)
    candidates, skipped = _list_candidates(folder)
    total = len(candidates)
    yield {"phase": "start", "total": total, "skipped": skipped}

    files, by_ext = [], {}
    for i, (path, name, ext, size) in enumerate(candidates, 1):
        yield {"phase": "file", "done": i, "total": total, "name": name}
        text = extract_text(path, ext)[:MAX_TEXT]
        files.append({
            "id": file_id(os.path.relpath(path, folder)),   # id stabile allo spostamento
            "rel": os.path.relpath(path, folder).replace("\\", "/"),
            "name": name, "ext": ext,
            "size": size, "mtime": os.path.getmtime(path),
            "text": text, "ntext": norm(text), "chars": len(text),
            "ocr": ext in IMG_EXT and bool(text),
        })
        by_ext[ext] = by_ext.get(ext, 0) + 1

    index = {
        "folder": folder,             # assoluto: solo per riferimento/display
        "folder_rel": folder_rel,     # relativo a base_dir se portabile
        "scanned_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "count": len(files), "skipped": skipped, "by_ext": by_ext,
        "ocr_available": OCR_AVAILABLE, "files": files,
    }
    if callable(index_path):          # salvataggio delegato (database)
        index_path(index)
    else:
        tmp = index_path + ".tmp"
        with io.open(tmp, "w", encoding="utf-8") as f:
            json.dump(index, f, ensure_ascii=False)
        os.replace(tmp, index_path)
    done = index_summary(index)
    done["phase"] = "done"
    yield done


def scan_folder(folder, index_path, base_dir=None):
    """Versione non-streaming: scandisce, salva e ritorna il riepilogo."""
    summary = None
    for ev in scan_folder_stream(folder, index_path, base_dir):
        if ev.get("phase") == "done":
            summary = {k: v for k, v in ev.items() if k != "phase"}
    return summary


def index_summary(index):
    out = {k: index.get(k) for k in
           ("folder", "scanned_at", "count", "skipped", "by_ext", "ocr_available")}
    out["folder_rel"] = index.get("folder_rel")
    out["portable"] = bool(index.get("folder_rel"))
    return out


def load_index(index_path):
    try:
        with io.open(index_path, "r", encoding="utf-8-sig") as f:
            return json.load(f)
    except Exception:
        return None


# --------------------------------------------------------------- ricerca ---
PHONE_QUERY = re.compile(r"^[\d\s.\-()+/]+$")   # query fatta solo di cifre e separatori
_SEP = r"[\s.\-()/]*"                            # separatori ammessi tra le cifre


def build_matcher(nq, mode):
    """Costruisce la regex di ricerca sul testo normalizzato (1:1 col testo
    originale). Se la query è un NUMERO, le cifre possono essere separate da
    spazi/trattini/punti: così '+23562798882' trova anche '+235 62 79 88 82'."""
    digits = re.sub(r"\D", "", nq)
    if len(digits) >= 5 and PHONE_QUERY.match(nq):
        core = _SEP.join(re.escape(d) for d in digits)
        pat = r"(?:\+" + _SEP + r")?" + core             # + iniziale opzionale
        if mode == "total":
            pat = r"(?<!\d)" + pat + r"(?!\d)"           # confini di numero
        return re.compile(pat)
    esc = re.escape(nq)
    if mode == "total":
        return re.compile(r"(?<!\w)" + esc + r"(?!\w)")
    return re.compile(esc)


def _snippets_from(text, matches, max_snips=3, radius=60):
    """Estratti di contesto attorno ai match. Le sostituzioni sono 1:1
    (newline/tab -> spazio) così gli offset restano validi; gli spazi multipli
    li accorpa il browser in fase di rendering."""
    out = []
    for m in matches[:max_snips]:
        i, j = m.start(), m.end()
        a = max(0, i - radius)
        b = min(len(text), j + radius)
        pre = "…" if a > 0 else ""
        post = "…" if b < len(text) else ""
        snip = text[a:b].replace("\n", " ").replace("\r", " ").replace("\t", " ")
        out.append({"text": pre + snip + post,
                    "at": i - a + len(pre),      # offset del match nello snippet
                    "len": j - i})               # lunghezza REALE del match
    return out


def search_index(index, query, mode="partial", max_files=60):
    """Cerca 'query' nell'indice. mode: 'partial' (sottostringa) o 'total'
    (parola intera). Ritorna i file che contengono il termine, con anteprime.
    Per i numeri di telefono i separatori vengono ignorati."""
    nq = norm(query).strip()
    if not index or not nq:
        return {"query": query, "mode": mode, "count": 0, "results": []}

    matcher = build_matcher(nq, mode)
    results = []
    for f in index.get("files", []):
        text = f.get("text") or ""
        # uso la normalizzata già salvata (veloce, allineata 1:1); la ricalcolo
        # solo se manca del tutto (indici vecchi senza il campo 'ntext').
        nt = f.get("ntext")
        if nt is None:
            nt = norm(text)
        if not nt:
            continue
        matches = list(matcher.finditer(nt))
        if not matches:
            continue
        results.append({
            "id": f["id"], "name": f["name"], "rel": f["rel"],
            "ext": f["ext"], "occ": len(matches), "ocr": f.get("ocr", False),
            "snippets": _snippets_from(text, matches),
        })
    results.sort(key=lambda r: -r["occ"])
    total = len(results)
    return {"query": query, "mode": mode, "count": total,
            "results": results[:max_files], "truncated": total > max_files}


_WORDSEP = r"[\s.\-,/()]+"        # separatore tra parole di una frase


def _term_core(w):
    """Frammento regex per una parola: sep-tollerante se è un numero."""
    nw = norm(w)
    digits = re.sub(r"\D", "", nw)
    if len(digits) >= 5 and PHONE_QUERY.match(nw):
        return r"(?:\+" + _SEP + r")?" + _SEP.join(re.escape(d) for d in digits)
    return re.escape(nw)


def _combo_regex(words, mode):
    """Regex di una combinazione: le parole devono comparire consecutive
    (separate da spazi/punteggiatura). Per una sola parola, in modo 'total'
    aggiunge i confini di parola."""
    core = _WORDSEP.join(_term_core(w) for w in words)
    if mode == "total" and len(words) == 1:
        core = r"(?<!\w)" + core + r"(?!\w)"
    return re.compile(core)


def search_multi(index, terms, mode="partial", max_files=80, max_words=4, max_combos=400):
    """Cerca nelle fonti TUTTE le combinazioni delle parole della chiave scheda:
    da 1 a max_words parole, in ogni ordine (permutazioni). Ritorna le
    combinazioni che hanno prodotto risultati (per i chip) e i file che le
    contengono. Un pre-filtro per parola tiene i tempi bassi anche su indici
    grandi: una combinazione si cerca solo nei file che contengono tutte le
    sue parole."""
    seen, words = set(), []
    for t in terms:
        t = str(t or "").strip()
        k = norm(t)
        if len(t) >= 2 and k and k not in seen:
            seen.add(k)
            words.append(t)
    if not index or not words:
        return {"terms": [], "results": [], "total": 0, "truncated": False}

    files = index.get("files", [])
    nts = []
    for f in files:
        nt = f.get("ntext")
        nts.append(nt if nt is not None else norm(f.get("text") or ""))

    # 1) pre-filtro: quali file contengono ciascuna parola
    word_rx = {w: _combo_regex([w], "partial") for w in words}
    word_files = {w: set() for w in words}
    for i, nt in enumerate(nts):
        if not nt:
            continue
        for w in words:
            if word_rx[w].search(nt):
                word_files[w].add(i)

    # 2) genera le combinazioni (permutazioni) di 1..max_words parole
    combos, seen_c = [], set()
    for r in range(1, min(max_words, len(words)) + 1):
        for perm in itertools.permutations(words, r):
            key = "\x00".join(norm(x) for x in perm)
            if key in seen_c:
                continue
            seen_c.add(key)
            combos.append(perm)
            if len(combos) >= max_combos:
                break
        if len(combos) >= max_combos:
            break

    # 3) cerca ogni combinazione come FRASE, solo nei file candidati
    combo_count = {}
    file_data = {}
    for perm in combos:
        cand = set.intersection(*[word_files[w] for w in perm]) if perm else set()
        if not cand:
            continue
        label = " ".join(perm)
        rx = _combo_regex(list(perm), mode)
        cnt = 0
        for i in cand:
            ms = list(rx.finditer(nts[i]))
            if not ms:
                continue
            cnt += 1
            fd = file_data.setdefault(i, {"matched": [], "occ": 0, "snippets": []})
            fd["matched"].append(label)
            fd["occ"] += len(ms)
            for sn in _snippets_from(files[i].get("text") or "", ms, max_snips=1):
                sn["term"] = label
                fd["snippets"].append(sn)
        if cnt:
            combo_count[label] = cnt

    # chip: le combinazioni con più parole in alto, poi per conteggio
    terms_sorted = sorted(
        ({"term": l, "count": c, "words": l.count(" ") + 1}
         for l, c in combo_count.items()),
        key=lambda x: (-x["words"], -x["count"], x["term"]))

    results = []
    for i, fd in file_data.items():
        f = files[i]
        # snippet: prima le combinazioni più lunghe (più specifiche)
        fd["snippets"].sort(key=lambda s: -(s["term"].count(" ") + 1))
        results.append({
            "id": f["id"], "name": f["name"], "rel": f["rel"], "ext": f["ext"],
            "ocr": f.get("ocr", False), "matched": fd["matched"],
            "occ": fd["occ"], "snippets": fd["snippets"][:8],
        })
    results.sort(key=lambda r: (-max((m.count(" ") + 1) for m in r["matched"]),
                                -len(r["matched"]), -r["occ"]))
    return {"terms": terms_sorted, "results": results[:max_files],
            "total": len(results), "truncated": len(results) > max_files}


def find_file(index, fid):
    for f in index.get("files", []):
        if f["id"] == fid:
            return f
    return None
